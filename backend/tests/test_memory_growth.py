"""The process grew until Render killed it. Two allocations did that.

HTTP/2: the shared Supabase client keeps a dead connection after a reset
(idle, not closed, not available) and opens another one next to it. A burst
of those resets — what autosave does to one client from many threads — stacks
connections whose buffers malloc_trim cannot release.

Shop File: clearing the template called Worksheet.cell() out to max_row, and
cell() creates missing cells. One leftover cell on a high row turned into a
dense grid of tens of thousands of cells on every report.
"""
import inspect
import os
import socket
import ssl
import subprocess
import tempfile
import threading

import h2.config
import h2.connection
import h2.events
import httpx
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import db
import excel
import main


def _cert():
    directory = tempfile.mkdtemp()
    cert = os.path.join(directory, "cert.pem")
    key = os.path.join(directory, "key.pem")
    subprocess.check_call(
        [
            "openssl", "req", "-x509", "-newkey", "rsa:2048",
            "-keyout", key, "-out", cert, "-days", "1", "-nodes",
            "-subj", "/CN=localhost",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return cert, key


def _serve_h1(conn):
    try:
        while True:
            data = b""
            while b"\r\n\r\n" not in data:
                chunk = conn.recv(65535)
                if not chunk:
                    return
                data += chunk
            body = b'[{"ok":true}]'
            conn.sendall(
                b"HTTP/1.1 200 OK\r\n"
                b"content-type: application/json\r\n"
                b"content-length: " + str(len(body)).encode() + b"\r\n"
                b"connection: keep-alive\r\n\r\n" + body
            )
    except Exception:
        pass
    finally:
        try:
            conn.close()
        except Exception:
            pass


def _serve_h2(conn):
    h2c = h2.connection.H2Connection(config=h2.config.H2Configuration(client_side=False))
    h2c.initiate_connection()
    try:
        conn.sendall(h2c.data_to_send())
        while True:
            data = conn.recv(65535)
            if not data:
                break
            for event in h2c.receive_data(data):
                if isinstance(event, h2.events.RequestReceived):
                    body = b'[{"ok":true}]'
                    h2c.send_headers(event.stream_id, [
                        (":status", "200"),
                        ("content-type", "application/json"),
                        ("content-length", str(len(body))),
                    ])
                    h2c.send_data(event.stream_id, body, end_stream=True)
            outgoing = h2c.data_to_send()
            if outgoing:
                conn.sendall(outgoing)
    except Exception:
        pass
    finally:
        try:
            conn.close()
        except Exception:
            pass


def _handle(raw, cert, key):
    ctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    ctx.load_cert_chain(cert, key)
    ctx.set_alpn_protocols(["h2", "http/1.1"])
    try:
        conn = ctx.wrap_socket(raw, server_side=True)
    except Exception:
        raw.close()
        return
    if conn.selected_alpn_protocol() == "h2":
        _serve_h2(conn)
    else:
        _serve_h1(conn)


@pytest.fixture(scope="module")
def tls_server():
    cert, key = _cert()
    srv = socket.socket()
    srv.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
    srv.bind(("127.0.0.1", 0))
    srv.listen(64)
    port = srv.getsockname()[1]

    def accept_loop():
        while True:
            try:
                client, _addr = srv.accept()
            except OSError:
                return
            threading.Thread(target=_handle, args=(client, cert, key), daemon=True).start()

    threading.Thread(target=accept_loop, daemon=True).start()
    yield f"https://127.0.0.1:{port}"
    srv.close()


def _kill_pooled_sockets(client):
    for conn in list(client._transport._pool.connections):
        inner = getattr(conn, "_connection", None)
        stream = getattr(inner, "_network_stream", None)
        if stream is not None:
            try:
                stream.close()
            except Exception:
                pass


def _open_unpatched(base_url, http2):
    """A client that still honors http2=True, which is what the SDK used to build."""
    client = httpx.Client.__new__(httpx.Client)
    db._HTTPX_CLIENT_INIT(
        client,
        base_url=base_url,
        verify=False,
        timeout=5,
        http2=http2,
    )
    return client


def test_http2_reset_stacks_dead_connections_and_the_app_client_does_not(tls_server):
    """Old behavior keeps one more dead HTTP/2 connection per reset.

    The client the app actually builds overwrites http2=True, speaks HTTP/1,
    and the pool does not grow across the same resets.
    """
    legacy = _open_unpatched(tls_server, http2=True)
    first = legacy.get("/")
    assert first.http_version == "HTTP/2"
    for _ in range(6):
        _kill_pooled_sockets(legacy)
        try:
            legacy.get("/")
        except Exception:
            pass
    dead = [
        conn for conn in legacy._transport._pool.connections
        if not conn.is_available()
    ]
    assert len(legacy._transport._pool.connections) >= 4
    assert len(dead) >= 3
    legacy.close()

    app_client = httpx.Client(base_url=tls_server, verify=False, timeout=5, http2=True)
    try:
        assert app_client._transport._pool._http2 is False
        negotiated = app_client.get("/")
        assert negotiated.http_version == "HTTP/1.1"
        for _ in range(6):
            _kill_pooled_sockets(app_client)
            try:
                app_client.get("/")
            except Exception:
                app_client.get("/")
        assert len(app_client._transport._pool.connections) <= 2
    finally:
        app_client.close()


def test_supabase_sessions_are_http1():
    postgrest = db.supabase_admin.postgrest.session
    auth = db.supabase_admin.auth._http_client
    assert postgrest._transport._pool._http2 is False
    assert auth._transport._pool._http2 is False
    assert postgrest._transport._pool._max_keepalive_connections <= 2


def test_keep_alive_does_not_kill_the_process():
    source = inspect.getsource(main.keep_alive)
    assert "os.kill" not in source
    assert "signal." not in source
    assert "RESTART_THRESHOLD" not in source


def _save(wb):
    import io
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_shop_file_does_not_materialize_rows_up_to_max_row():
    """A stray cell on row 2000 used to create 1999×40 cells. It must not."""
    raw = _save(_template_workbook())
    excel._template_cache["shopfile"] = raw
    try:
        output, filename = excel.generate_shop_file(
            [{
                "retailer_name": "Costco",
                "program": "RTL-ATT-EDM",
                "store_number": "1287",
                "city": "Portland",
                "state": "OR",
                "visit_date": "2026-10-01",
                "visit_time": "10:30:00",
                "reps_present": "Pass",
            }],
            "Kelsey",
        )
    finally:
        excel._template_cache.pop("shopfile", None)

    assert filename.startswith("Shop File Kelsey ")
    wb = load_workbook(output)
    ws = wb.active
    # Header row + the one visit row + the stray cell that was already there.
    # The old loop left ~80,000 cells (1999 rows × 40 columns).
    assert len(ws._cells) < 80
    assert ws.cell(2, 2).value == "Costco"
    assert ws.cell(2, 2).font.bold is True
    assert ws.cell(2, 2).font.name == "Courier New"
    assert ws.cell(2000, 1).value is None
    wb.close()


def _template_workbook():
    wb = Workbook()
    ws = wb.active
    for col in range(1, 6):
        ws.cell(1, col, f"Header {col}")
    styled = ws.cell(2, 2, "SAMPLE")
    styled.font = Font(name="Courier New", bold=True)
    ws.cell(2000, 1, "STRAY")
    return wb
